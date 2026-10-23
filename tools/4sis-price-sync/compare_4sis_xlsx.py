# Сверка цен 4SIS (с 23.10.2026; ночью из daily_4sis.sh): наши товары против фида поставщика по коду 1С, ссылки на 4sis.ru из обхода.
import datetime, io, re, sys, xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

D = sys.argv[1]; OUT = sys.argv[2]

def norm(s):
    s = s.lower().replace('ё', 'е')
    return ' '.join(re.sub(r'[^\w]+', ' ', s).split())

feed = {}
root = ET.parse(D + '/feed_now.xml').getroot()
cats = {c.get('id'): c.text for c in root.iter('category')}
byVendor = {}
for o in root.iter('offer'):
    t = lambda k: (o.findtext(k) or '').strip()
    price = float(t('price') or 0); old = float(t('oldprice') or 0)
    f = dict(code=o.get('id'), name=t('name'), price=price, old=old, regular=max(price, old), vendor=t('vendorCode'),
             avail='да' if o.get('available') == 'true' else 'нет', stock=t('stock'), cat=cats.get(t('categoryId'), ''))
    feed[f['code']] = f
    byVendor.setdefault(f['vendor'], []).append(f['code'])

site = {}
for fn in ('4sis.tsv', '4sis_parts.tsv'):
    for l in io.open(D + '/' + fn, encoding='utf-8'):
        c = l.rstrip('\n').split('\t')
        if len(c) > 2 and c[2]:
            site.setdefault(norm(c[2]), c[0])

ours = []
for l in io.open(D + '/ours_cmp.tsv', encoding='utf-8'):
    c = l.rstrip('\n').split('\t')
    if len(c) >= 8:
        ours.append(dict(id=c[0], active=c[1], code=c[2], name=c[3], url=c[4], price=float(c[5] or 0), vendor=c[6], section=c[7]))

rows, used = [], set()
for o in ours:
    f = feed.get(o['code']); how = 'код 1С'
    if not f and o['vendor'] and len(byVendor.get(o['vendor'], [])) == 1:
        f = feed[byVendor[o['vendor']][0]]; how = 'артикул'
    if f:
        used.add(f['code'])
    rows.append((o, f, how if f else ''))
for code, f in feed.items():
    if code not in used and f['name']:
        rows.append((None, f, ''))

def status(o, f):
    if not f:
        return 'нет в фиде 4SIS' + ('' if o['active'] == 'Y' else ' (у нас скрыт)')
    if not o:
        return 'нет у нас'
    st = '' if o['active'] == 'Y' else 'у нас скрыт; '
    if not f['regular']:
        return st + 'в фиде без цены'
    d = o['price'] - f['regular']
    if abs(d) < 1:
        return st + 'совпадает'
    return st + ('у нас дороже' if d > 0 else 'у нас дешевле')

head = ['Статус', 'Разница, ₽ (наша − 4SIS без акции)', 'Разница, %', 'Название у нас', 'Цена у нас, ₽', 'Цена 4SIS без акции, ₽',
        'Цена 4SIS по акции, ₽', 'Скидка 4SIS, %', 'Наличие у 4SIS', 'Остаток у 4SIS', 'Раздел у нас', 'Ссылка latitudo.ru',
        'Ссылка 4sis.ru', 'Название у 4SIS', 'Код 1С', 'Артикул производителя', 'ID у нас', 'На сайте', 'Сопоставлено по']
out = []
for o, f, how in rows:
    diff = pct = None
    if o and f and f['regular']:
        diff = o['price'] - f['regular']; pct = diff / f['regular']
    sale = f['price'] if f and f['old'] and f['old'] > f['price'] else None
    disc = (1 - f['price'] / f['old']) if sale else None
    name4 = f['name'] if f else ''
    out.append([status(o, f), diff, pct, o['name'] if o else '', o['price'] if o else None, f['regular'] if f else None, sale, disc,
                f['avail'] if f else '', int(f['stock']) if f and f['stock'].isdigit() else (f['stock'] if f else ''),
                o['section'] if o else '', o['url'] if o else '', site.get(norm(name4), '') if f else '', name4,
                (f['code'] if f else o['code']), (f['vendor'] if f else o['vendor']), int(o['id']) if o else None,
                ('да' if o['active'] == 'Y' else 'скрыт') if o else '', how])

order = {'у нас дешевле': 0, 'у нас дороже': 1}
out.sort(key=lambda r: (order.get(r[0], 2 if 'совпадает' not in r[0] else 3), -abs(r[1] or 0)))

wb = Workbook(); ws = wb.active; ws.title = 'Сверка цен 4SIS'
ws.append(head)
fills = {'у нас дешевле': 'FDE2E2', 'у нас дороже': 'FFF2CC', 'совпадает': 'E2F0D9'}
for i, r in enumerate(out, start=2):
    ws.append(r)
    for key, col in fills.items():
        if r[0].endswith(key):
            for c in range(1, 4):
                ws.cell(i, c).fill = PatternFill('solid', fgColor=col)
    for c in (2, 5, 6, 7, 10):
        ws.cell(i, c).number_format = '# ##0'
    for c in (3, 8):
        ws.cell(i, c).number_format = '0.0%'
    for c in (12, 13):
        v = ws.cell(i, c).value
        if v:
            ws.cell(i, c).hyperlink = v; ws.cell(i, c).font = Font(color='0563C1', underline='single')
for c, w in enumerate([24, 14, 10, 60, 13, 13, 13, 10, 10, 10, 24, 40, 40, 50, 14, 26, 9, 9, 12], start=1):
    ws.column_dimensions[get_column_letter(c)].width = w
for c in range(1, len(head) + 1):
    ws.cell(1, c).font = Font(bold=True); ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'E2'; ws.auto_filter.ref = ws.dimensions

cnt = {}
for r in out:
    cnt[r[0]] = cnt.get(r[0], 0) + 1
s = wb.create_sheet('Итоги', 0)
s.append(['Сверка цен 4SIS: latitudo.ru и фид поставщика — ' + datetime.date.today().strftime('%d.%m.%Y')]); s['A1'].font = Font(bold=True, size=14)
s.append(['Сопоставление по коду 1С (если код у 4SIS сменился — по артикулу производителя). Цена 4SIS — без акции; акционная — отдельной колонкой.'])
s.append(['У нас ставим цену 4SIS без акции, поэтому «совпадает» = наша цена равна их цене без акции.'])
s.append([])
s.append(['Статус', 'Товаров'])
s['A5'].font = Font(bold=True); s['B5'].font = Font(bold=True)
for k in sorted(cnt, key=lambda x: -cnt[x]):
    s.append([k, cnt[k]])
s.append([])
s.append(['Наших товаров 4SIS', len(ours)]); s.append(['из них на сайте', sum(o['active'] == 'Y' for o in ours)])
s.append(['Товаров в фиде 4SIS', len(feed)]); s.append(['из них по акции', sum(1 for f in feed.values() if f['old'] > f['price'])])
s.column_dimensions['A'].width = 45; s.column_dimensions['B'].width = 12
wb.save(OUT)
print(len(out), cnt)
